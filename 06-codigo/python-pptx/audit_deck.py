"""
Validador automatico de um .pptx conforme as regras do compendio (07-checklists/auditoria-de-arquivo.md).

Uso:
    python audit_deck.py deck.pptx [--audience tatico|estrategico|operacional|externo]
                                   [--expect-logo] [--fonts Calibri,Arial] [--json saida.json]

Saida: tabela slide / elemento / ocorrencia / severidade (erro | aviso) e codigo de retorno 1 se houver erro.
Nao substitui a inspecao visual: e a primeira peneira.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

EMU = 914400
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
MIN_FONT = {"estrategico": 12, "tatico": 12, "externo": 12, "operacional": 18}
FOOTER_MIN = 9
MAX_BULLETS = 8
MAX_TABLE_ROWS = 11  # cabecalho + 10


@dataclass
class Finding:
    slide: int
    element: str
    issue: str
    severity: str  # erro | aviso


def inches(v) -> float:
    return (v or 0) / EMU


def text_runs(shape):
    if not shape.has_text_frame:
        return []
    return [(p, r) for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()]


def audit(path: Path, audience: str = "tatico", expect_logo: bool = False, fonts: set[str] | None = None) -> list[Finding]:
    f: list[Finding] = []
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        return [Finding(0, "arquivo", f"nao abre com python-pptx: {e}", "erro")]

    if abs(inches(prs.slide_width) - SLIDE_W) > 0.05 or abs(inches(prs.slide_height) - SLIDE_H) > 0.05:
        f.append(Finding(0, "arquivo", f"tamanho {inches(prs.slide_width):.2f}x{inches(prs.slide_height):.2f} pol. (esperado 16:9 13,33x7,5)", "aviso"))

    min_font = MIN_FONT.get(audience, 12)
    fonts_used: Counter = Counter()
    logo_found = False

    for i, slide in enumerate(prs.slides, start=1):
        has_data = False
        has_source = False
        has_title = False
        dark_bg = False
        try:
            fill = slide.background.fill
            dark_bg = fill.type is not None and fill.fore_color.rgb is not None and sum(fill.fore_color.rgb) < 300
        except Exception:  # noqa: BLE001
            dark_bg = False
        for sh in slide.shapes:
            name = sh.name
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                logo_found = logo_found or (inches(sh.width) <= 2.0 and inches(sh.height) <= 0.8)
            # bounds
            if sh.left is not None:
                l, tp, w, h = inches(sh.left), inches(sh.top), inches(sh.width), inches(sh.height)
                full_bleed = l <= 0.01 and w >= SLIDE_W - 0.02
                if l < -0.01 or tp < -0.01 or l + w > SLIDE_W + 0.01 or tp + h > SLIDE_H + 0.01:
                    f.append(Finding(i, name, f"fora do slide ({l:.2f},{tp:.2f} {w:.2f}x{h:.2f})", "erro"))
                elif not full_bleed and (l < MARGIN - 0.05 or l + w > SLIDE_W - MARGIN + 0.05) and sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                    if not (sh.has_text_frame and not sh.text_frame.text.strip()):
                        f.append(Finding(i, name, f"fora da zona segura horizontal ({l:.2f} a {l + w:.2f})", "aviso"))
            # placeholders
            if sh.is_placeholder:
                if sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
                    has_title = bool(sh.has_text_frame and sh.text_frame.text.strip())
                    if has_title:
                        title = sh.text_frame.text.strip()
                        size = next((r.font.size.pt for _, r in text_runs(sh) if r.font.size), 28)
                        chars_per_line = max(20, int(inches(sh.width) * 72 / (size * 0.5)))
                        if len(title) > 2 * chars_per_line:
                            f.append(Finding(i, "titulo", f"titulo longo (~{len(title)} caracteres, mais de 2 linhas)", "aviso"))
                        if len(title.split()) < 3:
                            f.append(Finding(i, "titulo", f"titulo curto demais para ser uma conclusao: '{title}'", "aviso"))
                elif sh.has_text_frame and not sh.text_frame.text.strip():
                    f.append(Finding(i, name, "placeholder vazio", "erro"))
            # charts / tables
            if getattr(sh, "has_chart", False) and sh.has_chart:
                has_data = True
                ch = sh.chart
                if ch.has_title:
                    f.append(Finding(i, name, "grafico com titulo interno (o titulo do slide ja e a conclusao)", "aviso"))
                if len(list(ch.plots[0].series)) > 4:
                    f.append(Finding(i, name, "grafico com mais de 4 series", "aviso"))
            if getattr(sh, "has_table", False) and sh.has_table:
                has_data = True
                tbl = sh.table
                if len(tbl.rows) > MAX_TABLE_ROWS:
                    f.append(Finding(i, name, f"tabela com {len(tbl.rows) - 1} linhas (max 10 no corpo)", "aviso"))
                for r_i, row in enumerate(tbl.rows):
                    for c_i, cell in enumerate(row.cells):
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if r.font.size and r.font.size.pt < min_font and r.text.strip():
                                    f.append(Finding(i, f"{name} célula {r_i},{c_i}", f"fonte {r.font.size.pt:.0f} pt abaixo do minimo {min_font}", "erro"))
                                    break
                            else:
                                continue
                            break
            # text boxes
            runs = text_runs(sh)
            if runs:
                text = sh.text_frame.text
                if text.strip().lower().startswith("fonte:"):
                    has_source = True
                is_footer = inches(sh.top) >= 6.7
                sizes = [r.font.size.pt for _, r in runs if r.font.size]
                for _, r in runs:
                    if r.font.name:
                        fonts_used[r.font.name] += 1
                if sizes:
                    smallest = min(sizes)
                    limit = FOOTER_MIN if is_footer else min_font
                    if smallest < limit and not (is_footer and smallest >= FOOTER_MIN):
                        # tolerate small caps labels (kicker/header/labels) >= 9pt on non-operational decks
                        if not (smallest >= 9 and audience != "operacional" and len(text) <= 60):
                            f.append(Finding(i, name, f"fonte {smallest:.0f} pt abaixo do minimo {limit} pt", "erro" if smallest < 9 else "aviso"))
                    # overflow heuristic: chars * (size^2 * 0.55) vs box area (pt^2)
                    avg = sum(sizes) / len(sizes)
                    box_area = inches(sh.width) * 72 * inches(sh.height) * 72
                    needed = len(text) * avg * avg * 0.55
                    if box_area > 0 and needed > box_area * 1.25 and len(text) > 30:
                        f.append(Finding(i, name, f"possivel estouro de texto ({len(text)} caracteres a {avg:.0f} pt em {inches(sh.width):.1f}x{inches(sh.height):.1f} pol.)", "aviso"))
                n_par = len([p for p in sh.text_frame.paragraphs if p.text.strip()])
                if n_par > MAX_BULLETS and not is_footer:
                    f.append(Finding(i, name, f"{n_par} paragrafos/bullets em uma caixa (max {MAX_BULLETS})", "aviso"))
        if has_data and not has_source and not dark_bg:
            f.append(Finding(i, "rodape", "slide com grafico/tabela sem 'Fonte:' no rodape", "aviso"))
        if not has_title and not dark_bg and len(slide.shapes) > 2:
            f.append(Finding(i, "titulo", "slide de conteudo sem titulo no placeholder", "aviso"))

    if expect_logo and not logo_found:
        f.append(Finding(0, "logotipo", "nenhuma imagem pequena (logo) encontrada no deck", "aviso"))
    if fonts:
        bad = [fn for fn in fonts_used if fn not in fonts]
        if bad:
            f.append(Finding(0, "fontes", f"fontes fora da lista permitida: {sorted(bad)}", "aviso"))
    return f


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--audience", default="tatico", choices=list(MIN_FONT))
    ap.add_argument("--expect-logo", action="store_true")
    ap.add_argument("--fonts", help="lista separada por virgula de fontes permitidas")
    ap.add_argument("--json", help="salvar resultado em JSON")
    a = ap.parse_args(argv)
    fonts = {x.strip() for x in a.fonts.split(",")} if a.fonts else None
    findings = audit(Path(a.deck), a.audience, a.expect_logo, fonts)
    errors = [x for x in findings if x.severity == "erro"]
    print(f"{a.deck}: {len(findings)} ocorrencias ({len(errors)} erros, {len(findings) - len(errors)} avisos)")
    print(f"{'slide':>5}  {'sev':<5}  {'elemento':<28}  ocorrencia")
    for x in sorted(findings, key=lambda k: (k.slide, k.severity != "erro")):
        print(f"{x.slide:>5}  {x.severity:<5}  {x.element[:28]:<28}  {x.issue}")
    if a.json:
        Path(a.json).write_text(json.dumps([asdict(x) for x in findings], ensure_ascii=False, indent=2), encoding="utf-8")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
