"""
DeckBuilder: biblioteca base em python-pptx para gerar apresentacoes a partir de uma spec
(ver ../../schemas/deck-spec.schema.json) seguindo as regras do compendio:
grid 16:9 (01-fundamentos/grid-e-layout.md), paletas por papel de cor
(05-artefatos-visuais/paletas/) e modelos de slide (04-modelos-de-slides/).

Uso rapido:
    from deckbuilder import DeckBuilder
    db = DeckBuilder(palette="corporativa-azul")
    db.add_cover("Projeto Atlas", "Status semana 34")
    db.add_kpi_row("4 de 5 no verde", [{"label": "Prazo", "value": "-2 sem", "status": "warning"}])
    db.save("saida.pptx")

Ou, a partir de uma spec JSON:
    DeckBuilder.from_spec(spec_dict).save("saida.pptx")

Requisitos: python-pptx >= 1.0
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Iterable

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from fmt import auto as fmt_auto, fmt_number

ROOT = Path(__file__).resolve().parents[2]
PALETTES_DIR = ROOT / "05-artefatos-visuais" / "paletas"

# ---------------------------------------------------------------------------
# Grid (polegadas) - ver 01-fundamentos/grid-e-layout.md
# ---------------------------------------------------------------------------
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.5
CONTENT_W = 12.33
TITLE_BOX = (0.5, 0.4, 12.33, 0.9)
SUBTITLE_BOX = (0.5, 1.3, 12.33, 0.5)
CONTENT_TOP_NO_SUB = 1.5
CONTENT_TOP_SUB = 1.9
CONTENT_BOTTOM = 6.7
FOOTER_BOX = (0.5, 6.8, 12.33, 0.4)
GUTTER = 0.3
KICKER_BOX = (0.5, 0.42, 12.33, 0.28)
TITLE_BOX_K = (0.5, 0.68, 12.33, 0.9)
SUBTITLE_BOX_K = (0.5, 1.55, 12.33, 0.45)
CONTENT_TOP_K_NO_SUB = 1.8
CONTENT_TOP_K_SUB = 2.1
CALLOUT_H = 0.8
CALLOUT_STYLE = {  # kind -> (fill, bar, label_color, text_color)
    "conclusion": ("neutral_light", "primary", "primary", "neutral_dark"),
    "warning": ("FFF3D6", "warning", "warning", "neutral_dark"),
    "recommendation": ("EAF5EA", "success", "success", "neutral_dark"),
    "decision": ("primary_dark", None, "accent", "FFFFFF"),
}

# Tamanhos base de fonte (pt) - ver 01-fundamentos/tipografia.md
FS = {
    "cover_title": 42, "cover_sub": 22, "cover_meta": 13,
    "title": 28, "subtitle": 16, "body": 16, "small": 12, "footer": 10,
    "kpi_label": 12, "kpi_value": 40, "kpi_delta": 13,
    "big_number": 80, "section_number": 60, "section_title": 36,
    "table": 12, "chart": 11, "quote": 28, "closing": 38,
}

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}

STATUS_WORDS = {
    "success": {"verde", "ok", "concluido", "concluído", "baixa", "baixo", "sim", "no alvo", "success", "done"},
    "warning": {"amarelo", "atencao", "atenção", "media", "média", "medio", "médio", "em risco", "em andamento", "warning"},
    "danger": {"vermelho", "critico", "crítico", "alta", "alto", "atrasado", "atrasada", "danger", "nao", "não"},
}
STATUS_LABEL = {"success": "Verde", "warning": "Amarelo", "danger": "Vermelho", "neutral": "Neutro"}

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "area": XL_CHART_TYPE.AREA,
}
REFERENCE_SERIES = {"meta", "planejado", "orcado", "orçado", "orcamento", "orçamento", "budget", "target", "plan"}


def load_palette(name_or_path: str | dict | None) -> dict:
    """Carrega uma paleta pelo nome (05-artefatos-visuais/paletas/<nome>.json), por caminho ou dict."""
    if isinstance(name_or_path, dict):
        return name_or_path
    name = name_or_path or "corporativa-azul"
    p = Path(name)
    if not p.exists():
        p = PALETTES_DIR / f"{name}.json"
    if not p.exists():
        raise FileNotFoundError(f"Paleta nao encontrada: {name}")
    return json.loads(p.read_text(encoding="utf-8"))


def rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str.lstrip("#").upper())


def status_key(value: Any) -> str | None:
    """Converte texto livre (Verde, Alta, success...) em chave de status; None se nao reconhecido."""
    if value is None:
        return None
    v = str(value).strip().lower()
    for key, words in STATUS_WORDS.items():
        if v in words:
            return key
    return "neutral" if v in {"neutral", "neutro", "nao iniciado", "não iniciado", "-"} else None


class DeckBuilder:
    """Constroi um deck slide a slide, aplicando grid, paleta e tipografia do compendio."""

    def __init__(self, palette: str | dict | None = "corporativa-azul", font_scale: float | None = None,
                 confidentiality: str | None = None, logo: str | None = None, template: str | None = None,
                 brand: str | None = None, deck_name: str | None = None, date: str | None = None,
                 logo_light: str | None = None, logo_position: str = "footer",
                 template_layouts: dict | None = None, template_branding: bool = False):
        self.template = template
        self.template_branding = bool(template and template_branding)
        self.prs = Presentation(template) if template else Presentation()
        if template:
            self._clear_slides()
        else:
            self.prs.slide_width = Inches(SLIDE_W)
            self.prs.slide_height = Inches(SLIDE_H)
        self.pal = load_palette(palette)
        self.colors: dict = self.pal["colors"]
        fonts = self.pal.get("font", {})
        self.font_title = fonts.get("title", "Calibri")
        self.font_body = fonts.get("body", "Calibri")
        self.scale = float(font_scale if font_scale is not None else self.pal.get("font_scale", 1.0))
        self.confidentiality = confidentiality
        self.logo = logo if (logo and Path(logo).exists()) else (str(ROOT / logo) if logo and (ROOT / logo).exists() else None)
        self.logo_light = logo_light if (logo_light and Path(logo_light).exists()) else (str(ROOT / logo_light) if logo_light and (ROOT / logo_light).exists() else None)
        self.logo_position = logo_position or "footer"
        self.brand = brand
        self.deck_name = deck_name
        self.date = date
        self.page = 0
        self.bottom = CONTENT_BOTTOM
        self._layouts = self._pick_layouts(template_layouts or {})
        self._blank = self._layouts["blank"]
        self._title_only = self._layouts["content"]

    # ------------------------------------------------------------------ util
    def c(self, role: str) -> RGBColor:
        """Cor por papel ('primary', 'accent', ...) ou hex literal."""
        return rgb(self.colors.get(role, role))

    def chart_color(self, i: int) -> RGBColor:
        seq = self.colors.get("chart") or [self.colors["primary"], self.colors["secondary"], self.colors["accent"], self.colors["neutral_mid"]]
        return rgb(seq[i % len(seq)])

    def pt(self, key_or_size: str | float) -> Pt:
        size = FS[key_or_size] if isinstance(key_or_size, str) else key_or_size
        return Pt(size * self.scale)

    def _clear_slides(self):
        """Remove slides que vierem no arquivo de template."""
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            self.prs.part.drop_rel(sld.rId)
            lst.remove(sld)

    def _pick_layouts(self, mapping: dict) -> dict:
        """Escolhe layouts por nome (pt/en) ou por mapeamento explicito {cover, section, content, blank}."""
        layouts = list(self.prs.slide_layouts)
        names = [(l.name or "").lower() for l in layouts]
        aliases = {
            "cover": ["title slide", "slide de título", "slide de titulo", "capa", "title"],
            "section": ["section header", "cabeçalho da seção", "cabecalho da secao", "seção", "secao", "divis"],
            "content": ["title only", "somente título", "somente titulo", "apenas título", "apenas titulo", "título e conteúdo", "title and content"],
            "blank": ["blank", "em branco", "vazio"],
        }
        defaults = {"cover": 0, "section": 2, "content": 5, "blank": 6}
        out = {}
        for key in ("cover", "section", "content", "blank"):
            want = mapping.get(key)
            found = None
            if isinstance(want, int) and 0 <= want < len(layouts):
                found = layouts[want]
            elif isinstance(want, str):
                found = next((l for l, n in zip(layouts, names) if want.lower() == n or want.lower() in n), None)
            if found is None:
                for alias in aliases[key]:
                    found = next((l for l, n in zip(layouts, names) if alias in n), None)
                    if found:
                        break
            if found is None:
                idx = defaults[key] if defaults[key] < len(layouts) else len(layouts) - 1
                found = layouts[idx]
            out[key] = found
        return out

    def _new_slide(self, with_title: bool = False):
        slide = self.prs.slides.add_slide(self._title_only if with_title else self._blank)
        self.page += 1
        return slide

    def _fill_bg(self, slide, role: str):
        if self.template_branding:
            return
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = self.c(role)

    def _rect(self, slide, x, y, w, h, fill: str | None = "neutral_light", line: str | None = None,
              shape=MSO_SHAPE.RECTANGLE, line_width: float = 1.0):
        shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.c(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = self.c(line)
            shp.line.width = Pt(line_width)
        else:
            shp.line.fill.background()
        return shp

    def _text(self, slide, x, y, w, h, text: str = "", size: str | float = "body", bold: bool = False,
              color: str = "neutral_dark", align: str = "left", anchor=MSO_ANCHOR.TOP, italic: bool = False,
              font: str | None = None, wrap: bool = True):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        lines = str(text).split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[align]
            r = p.add_run()
            r.text = line
            self._style_run(r, size, bold, color, italic, font)
        return tb

    def _style_run(self, run, size="body", bold=False, color="neutral_dark", italic=False, font=None):
        run.font.size = self.pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font or self.font_body
        run.font.color.rgb = self.c(color)

    def _set_bullet(self, paragraph, char: str = "•", color: str = "primary", indent_in: float = 0.25, level: int = 0):
        pPr = paragraph._p.get_or_add_pPr()
        indent = Inches(indent_in)
        pPr.set("marL", str(int(indent * (level + 1))))
        pPr.set("indent", str(-int(indent)))
        for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buAutoNum"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(buClr, qn("a:srgbClr"))
        srgb.set("val", self.colors.get(color, color).lstrip("#").upper())
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", char)

    def _bullets(self, slide, x, y, w, h, items: Iterable[str], size: str | float = "body", color: str = "neutral_dark",
                 sub_bullets: dict | None = None, space_after: float = 8, bullet_color: str = "primary"):
        """Caixa de texto com bullets reais (a:buChar). Suporta **negrito** no inicio do item."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.05)
        first = True
        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(space_after)
            self._add_marked_runs(p, str(item), size, color)
            self._set_bullet(p, color=bullet_color)
            for sub in (sub_bullets or {}).get(str(idx), []):
                sp = tf.add_paragraph()
                sp.space_after = Pt(space_after / 2)
                self._add_marked_runs(sp, str(sub), (FS[size] if isinstance(size, str) else size) - 2, color)
                self._set_bullet(sp, char="–", color="neutral_mid", level=1)
        return tb

    def _add_marked_runs(self, paragraph, text: str, size, color):
        """Interpreta **trecho** como negrito."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if not part:
                continue
            r = paragraph.add_run()
            r.text = part
            self._style_run(r, size, bold=(i % 2 == 1), color=color)

    def _header(self, slide):
        """Linha discreta no topo: MARCA | DECK a esquerda; data ou logo a direita."""
        if self.template_branding:
            return
        left = "  |  ".join(v for v in (self.brand, self.deck_name) if v)
        if left:
            self._text(slide, 0.5, 0.1, 8.0, 0.25, left.upper(), size=9, bold=True, color="neutral_mid", anchor=MSO_ANCHOR.MIDDLE)
        if self.logo_position == "header" and self.logo:
            slide.shapes.add_picture(self.logo, Inches(11.6), Inches(0.1), height=Inches(0.32))
        elif self.date:
            self._text(slide, 8.5, 0.1, 4.33, 0.25, self.date, size=9, color="neutral_mid", align="right", anchor=MSO_ANCHOR.MIDDLE)

    def _place_logo(self, slide, x, y, width=None, height=None, dark=False):
        if self.template_branding:
            return
        src = (self.logo_light or self.logo) if dark else self.logo
        if src:
            kw = {"width": Inches(width)} if width else {"height": Inches(height)}
            slide.shapes.add_picture(src, Inches(x), Inches(y), **kw)
        elif self.brand:
            self._text(slide, x - 1.0, y, 2.5 if width is None else width, 0.4, self.brand.upper(), size=10, bold=True,
                       color="FFFFFF" if dark else "primary", align="right", anchor=MSO_ANCHOR.MIDDLE)

    def _title(self, slide, title: str | None, subtitle: str | None = None, kicker: str | None = None) -> float:
        """Escreve cabecalho, kicker, titulo (placeholder nativo) e subtitulo; devolve o top da area de conteudo."""
        self._header(slide)
        tbox, sbox = (TITLE_BOX_K, SUBTITLE_BOX_K) if kicker else (TITLE_BOX, SUBTITLE_BOX)
        if kicker:
            self._text(slide, *KICKER_BOX, str(kicker).upper(), size=10, bold=True, color="secondary", anchor=MSO_ANCHOR.BOTTOM)
        if title is not None:
            ph = slide.shapes.title
            if ph is None:
                self._text(slide, *tbox, title, size="title", bold=True, color="primary", font=self.font_title)
            elif self.template_branding:
                # posicao do grid, alinhamento a esquerda; fonte, cor e negrito ficam do template
                ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in tbox)
                tf = ph.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = title
                if r.font.size is None or r.font.size.pt > 32:
                    r.font.size = self.pt("title")
            else:
                ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in tbox)
                tf = ph.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.TOP
                tf.margin_left = tf.margin_right = Inches(0.05)
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = title
                self._style_run(r, "title", bold=True, color="primary", font=self.font_title)
        if subtitle:
            self._text(slide, *sbox, subtitle, size="subtitle", color="neutral_mid")
            return CONTENT_TOP_K_SUB if kicker else CONTENT_TOP_SUB
        return CONTENT_TOP_K_NO_SUB if kicker else CONTENT_TOP_NO_SUB

    def _footer(self, slide, source: str | None = None):
        x, y, w, h = FOOTER_BOX
        left = f"Fonte: {source}" if source else ""
        if self.confidentiality:
            left = f"{left}   |   {self.confidentiality}" if left else self.confidentiality
        if left:
            self._text(slide, x, y, w - 3.0, h, left, size="footer", color="neutral_mid", anchor=MSO_ANCHOR.MIDDLE)
        self._text(slide, x + w - 1.0, y, 1.0, h, str(self.page), size="footer", color="neutral_mid", align="right", anchor=MSO_ANCHOR.MIDDLE)
        if self.logo_position == "footer" and not self.template_branding:
            if self.logo:
                slide.shapes.add_picture(self.logo, Inches(10.6), Inches(6.83), height=Inches(0.3))
            elif self.brand:
                self._text(slide, 9.8, y, 2.0, h, self.brand.upper(), size=9, bold=True, color="primary", align="right", anchor=MSO_ANCHOR.MIDDLE)

    def _callout(self, slide, callout: dict):
        """Caixa rotulada no rodape da area de conteudo (ver 04-modelos-de-slides/callout.md)."""
        kind = callout.get("kind", "conclusion")
        fill, bar, label_color, text_color = CALLOUT_STYLE.get(kind, CALLOUT_STYLE["conclusion"])
        y = CONTENT_BOTTOM - CALLOUT_H - 0.05
        self._rect(slide, 0.5, y, CONTENT_W, CALLOUT_H, fill=fill)
        if bar:
            self._rect(slide, 0.5, y, 0.08, CALLOUT_H, fill=bar)
        label = callout.get("label")
        tx = 0.75
        if label:
            self._text(slide, 0.75, y, 2.3, CALLOUT_H, str(label).upper(), size=10, bold=True, color=label_color, anchor=MSO_ANCHOR.MIDDLE)
            tx = 3.05
        self._text(slide, tx, y, 12.83 - tx - 0.2, CALLOUT_H, callout.get("text", ""), size=14, bold=True, color=text_color, anchor=MSO_ANCHOR.MIDDLE)

    def _notes(self, slide, notes: str | None):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def _content_slide(self, spec_title, subtitle=None, source=None, notes=None, kicker=None, callout=None):
        slide = self._new_slide(with_title=True)
        for ph in list(slide.placeholders):
            if "TITLE" not in str(ph.placeholder_format.type):
                ph._element.getparent().remove(ph._element)
        top = self._title(slide, spec_title, subtitle, kicker)
        self._footer(slide, source)
        self._notes(slide, notes)
        self.bottom = CONTENT_BOTTOM
        if callout:
            self._callout(slide, callout)
            self.bottom = CONTENT_BOTTOM - CALLOUT_H - 0.2
        return slide, top

    @staticmethod
    def _columns(n: int, x0: float = MARGIN, total_w: float = CONTENT_W, gutter: float = GUTTER):
        w = (total_w - gutter * (n - 1)) / n
        return [(x0 + i * (w + gutter), w) for i in range(n)]

    # --------------------------------------------------------------- slides
    def add_cover(self, title: str, subtitle: str | None = None, author: str | None = None, date: str | None = None,
                  notes: str | None = None, kicker: str | None = None, thesis: str | None = None):
        if self.template_branding:
            s = self.prs.slides.add_slide(self._layouts["cover"])
            self.page += 1
            self._fill_placeholders(s, title, subtitle or (thesis or ""))
            meta = " | ".join(v for v in (author, date or self.date) if v)
            if meta:
                self._text(s, 0.8, 6.2, 11.7, 0.5, meta, size="cover_meta", color="neutral_mid")
            self._notes(s, notes)
            return s
        s = self._new_slide()
        self._fill_bg(s, "primary_dark")
        self._rect(s, 0, 6.9, SLIDE_W, 0.6, fill="accent")
        if self.brand:
            self._text(s, 0.8, 0.5, 6.0, 0.4, self.brand.upper(), size=11, bold=True, color="D9E1EA", anchor=MSO_ANCHOR.MIDDLE)
        if kicker:
            self._text(s, 0.8, 1.7, 11.7, 0.4, str(kicker).upper(), size=11, bold=True, color="accent", anchor=MSO_ANCHOR.BOTTOM)
        tw = 8.0 if thesis else 11.7
        self._text(s, 0.8, 2.1, tw, 1.9, title, size="cover_title", bold=True, color="FFFFFF", anchor=MSO_ANCHOR.BOTTOM, font=self.font_title)
        if subtitle:
            self._text(s, 0.8, 4.1, tw, 0.8, subtitle, size="cover_sub", color="D9E1EA")
        if thesis:
            self._rect(s, 9.2, 2.1, 3.63, 2.7, fill="FFFFFF", line=None)
            self._rect(s, 9.2, 2.1, 0.08, 2.7, fill="accent")
            self._text(s, 9.45, 2.25, 3.25, 0.4, "TESE CENTRAL", size=9, bold=True, color="neutral_mid")
            self._text(s, 9.45, 2.65, 3.25, 2.05, thesis, size=15, bold=True, color="primary")
        meta = " | ".join(v for v in (author, date or self.date, self.confidentiality) if v)
        if meta:
            self._text(s, 0.8, 6.2, 11.7, 0.5, meta, size="cover_meta", color="BFC9D4")
        self._place_logo(s, 11.3, 0.5, width=1.5, dark=True)
        self._notes(s, notes)
        return s

    def _fill_placeholders(self, slide, title: str, subtitle: str = ""):
        """Preenche titulo e subtitulo/corpo do layout do template; remove placeholders vazios."""
        for ph in list(slide.placeholders):
            kind = str(ph.placeholder_format.type)
            if "TITLE" in kind and "SUB" not in kind:
                ph.text_frame.text = title
            elif ("SUBTITLE" in kind or "BODY" in kind) and subtitle:
                ph.text_frame.text = subtitle
                subtitle = ""
            else:
                ph._element.getparent().remove(ph._element)

    def add_section(self, title: str, number: str | None = None, subtitle: str | None = None, notes: str | None = None):
        if self.template_branding:
            s = self.prs.slides.add_slide(self._layouts["section"])
            self.page += 1
            self._fill_placeholders(s, f"{number}. {title}" if number else title, subtitle or "")
            self._notes(s, notes)
            return s
        s = self._new_slide()
        self._fill_bg(s, "primary")
        if number:
            self._text(s, 0.8, 2.2, 2.0, 1.0, str(number), size="section_number", bold=True, color="accent", font=self.font_title)
        self._text(s, 0.8, 3.3, 11.7, 1.2, title, size="section_title", bold=True, color="FFFFFF", font=self.font_title)
        self._place_logo(s, 11.6, 0.5, width=1.2, dark=True)
        if subtitle:
            self._text(s, 0.8, 4.5, 11.7, 0.8, subtitle, size=18, color="D9E1EA")
        self._notes(s, notes)
        return s

    def add_agenda(self, items: list[str], title: str = "Agenda", durations: list[str] | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, notes=notes, kicker=kicker, callout=callout)
        step = min(0.75, (self.bottom - top - 0.2) / max(len(items), 1))
        for i, item in enumerate(items):
            y = top + 0.2 + i * step
            circ = self._rect(s, 0.5, y, 0.5, 0.5, fill="primary", shape=MSO_SHAPE.OVAL)
            circ.text_frame.text = str(i + 1)
            circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            circ.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._style_run(circ.text_frame.paragraphs[0].runs[0], 14, bold=True, color="FFFFFF")
            self._text(s, 1.2, y, 7.5, 0.5, item, size=20, anchor=MSO_ANCHOR.MIDDLE)
            if durations and i < len(durations):
                self._text(s, 8.8, y, 4.0, 0.5, durations[i], size=16, color="neutral_mid", anchor=MSO_ANCHOR.MIDDLE)
        return s

    def add_executive_summary(self, headline: str, points: list[str], ask: str | None = None, title: str = "Sumário executivo",
                              source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, source=source, notes=notes, kicker=kicker, callout=callout)
        self._rect(s, 0.5, top, 12.33, 1.1, fill="neutral_light")
        self._rect(s, 0.5, top, 0.08, 1.1, fill="accent")
        self._text(s, 0.75, top, 11.9, 1.1, headline, size=22, bold=True, color="primary", anchor=MSO_ANCHOR.MIDDLE)
        y = top + 1.4
        row_h = min(0.9, (self.bottom - y - (0.9 if ask else 0)) / max(len(points), 1))
        for i, ptxt in enumerate(points):
            yy = y + i * row_h
            circ = self._rect(s, 0.6, yy + 0.05, 0.45, 0.45, fill="primary", shape=MSO_SHAPE.OVAL)
            circ.text_frame.text = str(i + 1)
            circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            circ.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._style_run(circ.text_frame.paragraphs[0].runs[0], 13, bold=True, color="FFFFFF")
            self._text(s, 1.25, yy, 11.5, row_h, ptxt, size=17, anchor=MSO_ANCHOR.TOP)
        if ask:
            ya = self.bottom - 0.8
            self._rect(s, 0.5, ya, 12.33, 0.75, fill="FFFFFF", line="accent", line_width=1.5)
            tb = self._text(s, 0.7, ya, 12.0, 0.75, "", anchor=MSO_ANCHOR.MIDDLE)
            p = tb.text_frame.paragraphs[0]
            r1 = p.add_run(); r1.text = "Pedimos:  "; self._style_run(r1, 16, bold=True, color="accent")
            r2 = p.add_run(); r2.text = ask; self._style_run(r2, 16, bold=True, color="neutral_dark")
        return s

    def add_kpi_row(self, title: str, kpis: list[dict], subtitle: str | None = None, bullets: list[str] | None = None,
                    source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = max(2, min(len(kpis), 6))
        cols = self._columns(n)
        has_desc = any(k.get("description") for k in kpis)
        card_h = (2.9 if has_desc else 2.6) if n <= 4 else (2.9 if has_desc else 2.4)
        value_size = {2: 48, 3: 44, 4: 40, 5: 34, 6: 30}[n]
        y = top + 0.2
        for (x, w), k in zip(cols, kpis):
            st = k.get("status", "neutral")
            bar = st if st != "neutral" else "neutral_mid"
            self._rect(s, x, y, w, card_h, fill="neutral_light")
            self._rect(s, x, y, w, 0.08, fill=bar)
            self._text(s, x + 0.15, y + 0.25, w - 0.3, 0.4, str(k.get("label", "")).upper(), size="kpi_label", color="neutral_mid")
            val = fmt_auto(k.get("value", ""), k.get("format"))
            vsize = value_size if len(val) <= 6 else max(20, int(value_size * 6 / len(val) * 1.15))
            self._text(s, x + 0.15, y + 0.65, w - 0.3, 1.1, val, size=vsize, bold=True, color="neutral_dark", anchor=MSO_ANCHOR.MIDDLE)
            if k.get("delta"):
                self._text(s, x + 0.15, y + 1.75, w - 0.3, 0.45, str(k["delta"]), size="kpi_delta", bold=True, color=bar)
            if k.get("description"):
                self._text(s, x + 0.15, y + 2.15, w - 0.3, 0.7, str(k["description"]), size=10, color="neutral_mid")
        if bullets:
            self._bullets(s, 0.5, y + card_h + 0.3, 12.33, self.bottom - (y + card_h + 0.3), bullets, size=15)
        return s

    def add_big_number(self, title: str, value: str, label: str | None = None, context: str | None = None,
                       subtitle: str | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        self._text(s, 0.5, top + 0.4, 12.33, 2.0, value, size="big_number", bold=True, color="primary", align="center", anchor=MSO_ANCHOR.MIDDLE, font=self.font_title)
        if label:
            self._text(s, 0.5, top + 2.5, 12.33, 0.6, label, size=20, color="neutral_mid", align="center")
        if context:
            self._text(s, 1.5, top + 3.3, 10.33, 1.4, context, size=16, align="center")
        return s

    def add_bullets(self, title: str, bullets: list[str], subtitle: str | None = None, sub_bullets: dict | None = None,
                    source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        self._bullets(s, 0.5, top + 0.1, 12.33, self.bottom - top - 0.1, bullets, size="body", sub_bullets=sub_bullets, space_after=10)
        return s

    def add_two_column(self, title: str, left: dict, right: dict, subtitle: str | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        for (x, w), col in zip(self._columns(2), (left, right)):
            st = col.get("status")
            head_color = st if st in ("success", "warning", "danger") else "primary"
            self._text(s, x, top + 0.1, w, 0.6, col.get("heading", ""), size=18, bold=True, color=head_color, anchor=MSO_ANCHOR.BOTTOM)
            self._rect(s, x, top + 0.75, w, 0.03, fill=head_color)
            self._bullets(s, x, top + 0.95, w, self.bottom - top - 0.95, col.get("bullets", []), size=15, bullet_color=head_color)
        return s

    def add_comparison(self, title: str, options: list[dict], subtitle: str | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = max(2, min(len(options), 4))
        cols = self._columns(n)
        card_h = self.bottom - top - 0.2
        text_size = 15 if n <= 3 else 13
        for (x, w), opt in zip(cols, options):
            rec = bool(opt.get("recommended"))
            self._rect(s, x, top + 0.1, w, card_h, fill="neutral_light", line="accent" if rec else None, line_width=2.0)
            self._rect(s, x, top + 0.1, w, 0.7, fill="primary")
            self._text(s, x, top + 0.1, w, 0.7, opt.get("name", ""), size=17 if n <= 3 else 15, bold=True, color="FFFFFF", align="center", anchor=MSO_ANCHOR.MIDDLE)
            if rec:
                self._rect(s, x + w - 1.6, top - 0.1, 1.5, 0.35, fill="accent")
                self._text(s, x + w - 1.6, top - 0.1, 1.5, 0.35, "Recomendada", size=10, bold=True, color="FFFFFF", align="center", anchor=MSO_ANCHOR.MIDDLE)
            self._bullets(s, x + 0.15, top + 1.0, w - 0.3, card_h - 1.0, opt.get("points", []), size=text_size)
        return s

    def add_table(self, title: str, columns: list[str], rows: list[list], max_rows: int = 10, paginate: bool = True, **kw):
        """Tabela; acima de max_rows linhas, divide em varios slides '(i/n)' (anexos)."""
        if not paginate or len(rows) <= max_rows:
            return self._add_table_page(title, columns, rows, **kw)
        chunks = [rows[i:i + max_rows] for i in range(0, len(rows), max_rows)]
        last = None
        for i, chunk in enumerate(chunks, start=1):
            kw2 = dict(kw)
            if i > 1:
                kw2["total_row"] = False
                kw2.pop("highlight_rows", None)
            elif kw.get("total_row"):
                kw2["total_row"] = False
            last = self._add_table_page(f"{title} ({i}/{len(chunks)})", columns, chunk, **kw2)
        return last

    def _add_table_page(self, title: str, columns: list[str], rows: list[list], subtitle: str | None = None, align: list[str] | None = None,
                  col_widths: list[float] | None = None, total_row: bool = False, highlight_rows: list[int] | None = None,
                  status_columns: list[int] | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None, font_size: float | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n_rows, n_cols = len(rows) + 1, len(columns)
        row_h = min(0.5, (self.bottom - top - 0.1) / n_rows)
        shape = s.shapes.add_table(n_rows, n_cols, Inches(0.5), Inches(top + 0.1), Inches(CONTENT_W), Inches(row_h * n_rows))
        tbl = shape.table
        tbl.first_row = True
        tbl.horz_banding = False
        if col_widths:
            for i, cw in enumerate(col_widths[:n_cols]):
                tbl.columns[i].width = Inches(cw)
        else:
            for i in range(n_cols):
                tbl.columns[i].width = Inches(CONTENT_W / n_cols)
        fsize = font_size or FS["table"]
        highlight_rows = set(highlight_rows or [])
        status_columns = set(status_columns or [])

        def fill_cell(cell, text, fill_role, color, bold=False, al="left", size=fsize):
            cell.fill.solid()
            cell.fill.fore_color.rgb = self.c(fill_role)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN[al]
            r = p.add_run()
            r.text = "" if text is None else str(text)
            self._style_run(r, size, bold=bold, color=color)

        for j, col in enumerate(columns):
            al = "center" if j in status_columns else (align[j] if align and j < len(align) else "left")
            fill_cell(tbl.cell(0, j), col, "primary", "FFFFFF", bold=True, al=al)
        for i, row in enumerate(rows, start=1):
            is_total = total_row and i == len(rows)
            base_fill = "neutral_light" if i % 2 == 0 else "FFFFFF"
            if (i - 1) in highlight_rows:
                base_fill = "FFF3D6"
            for j in range(n_cols):
                val = row[j] if j < len(row) else ""
                al = align[j] if align and j < len(align) else ("right" if isinstance(val, (int, float)) else "left")
                if j in status_columns:
                    st = status_key(val)
                    if st and st != "neutral":
                        fill_cell(tbl.cell(i, j), val, st, "FFFFFF", bold=True, al="center")
                        continue
                    al = "center"
                fill_cell(tbl.cell(i, j), val, base_fill, "neutral_dark", bold=is_total or (i - 1) in highlight_rows, al=al)
        return s

    def add_chart(self, title: str, chart_type: str, categories: list[str], series: list[dict], subtitle: str | None = None,
                  highlight_index: int | None = None, number_format: str | None = None, show_labels: bool | None = None,
                  commentary: list[str] | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        xl_type = CHART_TYPES[chart_type]
        data = CategoryChartData()
        data.categories = categories
        for ser in series:
            data.add_series(ser["name"], ser["values"])
        if commentary:
            (cx, cw), (tx, tw) = self._columns(2)
            cw = 8.0; tx = 8.8; tw = 4.03
        else:
            cx, cw, tx, tw = 0.5, CONTENT_W, None, None
        h = self.bottom - top - 0.1
        gframe = s.shapes.add_chart(xl_type, Inches(cx), Inches(top + 0.1), Inches(cw), Inches(h), data)
        chart = gframe.chart
        self._set_alt(gframe, title)
        chart.has_title = False
        chart.font.size = self.pt("chart")
        chart.font.name = self.font_body
        chart.font.color.rgb = self.c("neutral_dark")
        multi = len(series) > 1
        is_pie = chart_type in ("pie", "doughnut")
        chart.has_legend = multi or is_pie
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = self.pt("chart")
        if not is_pie:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.c("neutral_light")
            va.format.line.fill.background()
            va.tick_labels.font.size = self.pt("chart")
            va.tick_labels.font.color.rgb = self.c("neutral_mid")
            ca = chart.category_axis
            ca.tick_labels.font.size = self.pt("chart")
            ca.tick_labels.font.color.rgb = self.c("neutral_mid")
            ca.format.line.color.rgb = self.c("neutral_mid")
            ca.has_major_gridlines = False
            if chart_type in ("bar", "stacked_bar"):
                ca.reverse_order = True  # primeira categoria no topo (leitura natural do ranking)
                # eixo de valores cruza o de categorias no maximo (embaixo, ja que esta invertido)
                crosses = va._element.find(qn("c:crosses"))
                if crosses is not None:
                    crosses.set("val", "max")
        plot = chart.plots[0]
        labels_on = show_labels if show_labels is not None else ((not multi and len(categories) <= 12) or is_pie)
        if labels_on:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = self.pt("chart")
            dl.font.color.rgb = self.c("neutral_dark")
            if number_format:
                dl.number_format = number_format
                dl.number_format_is_linked = False
            if is_pie:
                dl.show_percentage = True
                dl.show_value = False
        if chart_type in ("bar", "column", "stacked_bar", "stacked_column"):
            plot.gap_width = 60
            if multi:
                plot.overlap = -10
        if is_pie:
            plot.vary_by_categories = True
            for idx, point in enumerate(plot.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self.chart_color(idx)
        else:
            for i, ser in enumerate(plot.series):
                is_ref = str(ser.name).strip().lower() in REFERENCE_SERIES
                color = self.c("neutral_mid") if (is_ref and multi) else self.chart_color(i)
                if chart_type in ("line",):
                    ser.smooth = False
                    ser.format.line.color.rgb = color
                    ser.format.line.width = Pt(2.5)
                    if is_ref and multi:
                        ser.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = color
                    ser.marker.format.line.color.rgb = color
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = color
                    ser.format.line.fill.background()
                    if highlight_index is not None and not multi and chart_type in ("bar", "column"):
                        pt_ = ser.points[highlight_index]
                        pt_.format.fill.solid()
                        pt_.format.fill.fore_color.rgb = self.c("accent")
        if commentary:
            self._bullets(s, tx, top + 0.2, tw, h - 0.2, commentary, size=14)
        return s

    def add_timeline(self, title: str, milestones: list[dict], subtitle: str | None = None, today_index: int | None = None,
                     source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = len(milestones)
        y_line = top + (self.bottom - top) / 2
        self._rect(s, 0.8, y_line - 0.03, 11.7, 0.06, fill="neutral_mid")
        xs = [1.3 + i * (10.7 / max(n - 1, 1)) for i in range(n)]
        for i, (x, m) in enumerate(zip(xs, milestones)):
            st = m.get("status", "neutral")
            color = st if st != "neutral" else "primary"
            self._rect(s, x - 0.2, y_line - 0.2, 0.4, 0.4, fill=color, line="FFFFFF", shape=MSO_SHAPE.OVAL, line_width=1.5)
            above = (i % 2 == 0) or n <= 6
            if above or n <= 6:
                self._text(s, x - 0.9, y_line - 1.0, 1.8, 0.4, m.get("date", ""), size=13, color="neutral_mid", align="center", anchor=MSO_ANCHOR.BOTTOM)
                self._text(s, x - 0.9, y_line + 0.35, 1.8, 1.2, m.get("label", ""), size=13, align="center")
            else:
                self._text(s, x - 0.9, y_line + 0.35, 1.8, 0.4, m.get("date", ""), size=13, color="neutral_mid", align="center")
                self._text(s, x - 0.9, y_line - 1.55, 1.8, 1.2, m.get("label", ""), size=13, align="center", anchor=MSO_ANCHOR.BOTTOM)
        if today_index is not None and 0 <= today_index < n:
            xt = xs[today_index] + (10.7 / max(n - 1, 1)) * 0.5 if today_index < n - 1 else xs[today_index] + 0.3
            ln = self._rect(s, xt, y_line - 0.7, 0.03, 1.4, fill="accent")
            self._text(s, xt - 0.5, y_line - 1.05, 1.0, 0.3, "Hoje", size=11, bold=True, color="accent", align="center")
        legend = [("success", "Concluído"), ("warning", "Em risco"), ("danger", "Atrasado"), ("primary", "Planejado")]
        used = {m.get("status", "neutral") for m in milestones}
        lx = 0.5
        for role, lab in legend:
            key = "neutral" if role == "primary" else role
            if key not in used:
                continue
            self._rect(s, lx, self.bottom - 0.3, 0.2, 0.2, fill=role, shape=MSO_SHAPE.OVAL)
            self._text(s, lx + 0.25, self.bottom - 0.35, 1.5, 0.3, lab, size=11, color="neutral_mid", anchor=MSO_ANCHOR.MIDDLE)
            lx += 1.7
        return s

    def add_process(self, title: str, steps: list[str], descriptions: list[str] | None = None, current_index: int | None = None, metrics: list[str] | None = None,
                    subtitle: str | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = len(steps)
        w = CONTENT_W / n
        y = top + 0.6
        for i, step in enumerate(steps):
            x = 0.5 + i * w
            shape = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
            fill = "accent" if current_index == i else "primary"
            shp = self._rect(s, x, y, w - 0.1, 1.4, fill=fill, shape=shape)
            tf = shp.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.35 if i else 0.15)
            tf.margin_right = Inches(0.3)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run(); r.text = f"{i + 1}. {step}"
            self._style_run(r, 14 if n <= 4 else 12, bold=True, color="FFFFFF")
            dy = y + 1.6
            if metrics and i < len(metrics) and metrics[i]:
                self._text(s, x + 0.1, dy, w - 0.3, 0.6, str(metrics[i]), size=26, bold=True, color="accent" if current_index == i else "primary")
                dy += 0.65
            if descriptions and i < len(descriptions) and descriptions[i]:
                self._text(s, x + 0.1, dy, w - 0.3, 1.4, descriptions[i], size=13, color="neutral_dark")
        return s

    def add_matrix_2x2(self, title: str, quadrants: dict, x_label: str | None = None, y_label: str | None = None, highlight: str | None = None,
                       subtitle: str | None = None, source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        mx, my, mw, mh = 1.6, top + 0.1, 9.0, self.bottom - top - 0.6
        qw, qh = mw / 2 - 0.05, mh / 2 - 0.05
        pos = {"tl": (mx, my), "tr": (mx + qw + 0.1, my), "bl": (mx, my + qh + 0.1), "br": (mx + qw + 0.1, my + qh + 0.1)}
        for key, (qx, qy) in pos.items():
            q = quadrants.get(key, {}) or {}
            self._rect(s, qx, qy, qw, qh, fill="FFF3D6" if highlight == key else "neutral_light")
            if q.get("label"):
                self._text(s, qx + 0.1, qy + 0.05, qw - 0.2, 0.35, q["label"].upper(), size=11, bold=True, color="neutral_mid")
            self._bullets(s, qx + 0.1, qy + 0.45, qw - 0.2, qh - 0.5, q.get("items", []), size=13, space_after=4)
        if x_label:
            self._text(s, mx, my + mh + 0.05, mw, 0.4, f"{x_label}  →", size=13, bold=True, color="neutral_mid", align="center")
        if y_label:
            tb = self._text(s, mx - 1.1, my + mh / 2 - 0.2, 1.0, 0.4, f"{y_label}  →", size=13, bold=True, color="neutral_mid", align="center")
            tb.rotation = 270
        return s

    def add_action_plan(self, title: str, actions: list[dict], decision: str | None = None, headers: list[str] | None = None, subtitle: str | None = None,
                        source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n_rows = len(actions) + 1
        avail = self.bottom - top - 0.1 - (0.95 if decision else 0)
        row_h = min(0.5, avail / n_rows)
        widths = [6.5, 2.0, 1.5, 2.33]
        shape = s.shapes.add_table(n_rows, 4, Inches(0.5), Inches(top + 0.1), Inches(CONTENT_W), Inches(row_h * n_rows))
        tbl = shape.table
        tbl.horz_banding = False
        for i, wdt in enumerate(widths):
            tbl.columns[i].width = Inches(wdt)

        def cell(cell_, text, fill, color, bold=False, al="left"):
            cell_.fill.solid(); cell_.fill.fore_color.rgb = self.c(fill)
            cell_.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell_.margin_left = cell_.margin_right = Inches(0.08)
            p = cell_.text_frame.paragraphs[0]; p.alignment = ALIGN[al]
            r = p.add_run(); r.text = str(text); self._style_run(r, "table", bold=bold, color=color)

        for j, h in enumerate(headers or ["Ação", "Dono", "Prazo", "Status"]):
            cell(tbl.cell(0, j), h, "primary", "FFFFFF", bold=True, al="center" if j == 3 else "left")
        for i, a in enumerate(actions, start=1):
            base = "neutral_light" if i % 2 == 0 else "FFFFFF"
            cell(tbl.cell(i, 0), a.get("action", "") or (f"{i:02d}" if headers else ""), base, "neutral_dark")
            cell(tbl.cell(i, 1), a.get("owner", ""), base, "neutral_dark")
            cell(tbl.cell(i, 2), a.get("due", ""), base, "neutral_dark", al="center")
            st = a.get("status", "neutral")
            labels = {"success": "Concluída", "warning": "Em andamento", "danger": "Atrasada", "neutral": "Não iniciada"}
            if headers:
                cell(tbl.cell(i, 3), a.get("note", "") if st in labels or not st else st, base, "neutral_dark")
            elif st == "neutral":
                cell(tbl.cell(i, 3), labels[st], base, "neutral_mid", al="center")
            elif st in labels:
                cell(tbl.cell(i, 3), labels[st], st, "FFFFFF", bold=True, al="center")
            else:
                cell(tbl.cell(i, 3), st, base, "neutral_dark", al="center")
        if decision:
            yd = self.bottom - 0.8
            self._rect(s, 0.5, yd, 12.33, 0.75, fill="FFF3D6")
            self._rect(s, 0.5, yd, 0.08, 0.75, fill="accent")
            tb = self._text(s, 0.75, yd, 11.9, 0.75, "", anchor=MSO_ANCHOR.MIDDLE)
            p = tb.text_frame.paragraphs[0]
            r1 = p.add_run(); r1.text = "Decisão pedida:  "; self._style_run(r1, 16, bold=True, color="accent")
            r2 = p.add_run(); r2.text = decision; self._style_run(r2, 16, bold=True, color="neutral_dark")
        return s

    def add_takeaways(self, title: str, items: list[dict], subtitle: str | None = None, source: str | None = None,
                      notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = max(2, min(len(items), 4))
        cols = self._columns(n)
        card_h = self.bottom - top - 0.3
        head_size, text_size = (18, 14) if n <= 3 else (15, 12)
        for i, ((x, w), it) in enumerate(zip(cols, items)):
            color = self.chart_color(i)
            self._rect(s, x, top + 0.2, w, card_h, fill="neutral_light")
            bar = self._rect(s, x, top + 0.2, w, 0.06, fill="FFFFFF")
            bar.fill.fore_color.rgb = color
            if it.get("tag"):
                pill = self._rect(s, x + 0.25, top + 0.45, min(2.4, w - 0.5), 0.32, fill="FFFFFF")
                pill.fill.fore_color.rgb = color
                self._text(s, x + 0.25, top + 0.45, min(2.4, w - 0.5), 0.32, str(it["tag"]).upper(), size=9, bold=True, color="FFFFFF", align="center", anchor=MSO_ANCHOR.MIDDLE)
            else:
                tb = self._text(s, x + 0.25, top + 0.4, 1.5, 0.5, f"{i + 1:02d}", size=22, bold=True)
                tb.text_frame.paragraphs[0].runs[0].font.color.rgb = color
            self._text(s, x + 0.25, top + 1.0, w - 0.5, 0.75, it.get("heading", ""), size=head_size, bold=True, anchor=MSO_ANCHOR.TOP)
            self._text(s, x + 0.25, top + 1.75, w - 0.5, card_h - 1.7, it.get("text", ""), size=text_size)
        return s

    def add_progress_bars(self, title: str, items: list[dict], subtitle: str | None = None, max_value: float | None = None,
                          columns: int = 1, highlight_index: int | None = None, source: str | None = None,
                          notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        vmax = float(max_value or 100)
        ncol = 2 if columns == 2 else 1
        per_col = (len(items) + ncol - 1) // ncol
        row_h = min(0.55, (self.bottom - top - 0.2) / max(per_col, 1))
        for i, it in enumerate(items):
            c, r = divmod(i, per_col)
            x0 = 0.5 + c * 6.32
            y = top + 0.15 + r * row_h
            if ncol == 1:
                lw, bx, bw, vx, vw = 3.0, 3.6, 7.0, 10.7, 2.13
            else:
                lw, bx, bw, vx, vw = 1.9, x0 + 2.0, 3.0, x0 + 5.1, 0.9
                bx, vx = x0 + 2.0, x0 + 5.1
            st = it.get("status")
            color = st if st in ("success", "warning", "danger") else ("accent" if highlight_index == i else "primary")
            self._text(s, x0, y, lw, 0.45, str(it.get("label", "")), size=14 if ncol == 1 else 12, anchor=MSO_ANCHOR.MIDDLE)
            self._rect(s, bx, y + 0.1, bw, 0.25, fill="neutral_light")
            frac = min(1.0, max(0.0, float(it.get("value", 0)) / vmax))
            if frac > 0:
                self._rect(s, bx, y + 0.1, bw * frac, 0.25, fill=color)
            self._text(s, vx, y, vw, 0.45, str(it.get("display", it.get("value", ""))), size=13 if ncol == 1 else 11, bold=True,
                       color=color if st else "neutral_dark", anchor=MSO_ANCHOR.MIDDLE)
        return s

    def _resolve(self, path: str) -> str:
        for cand in (Path(path), ROOT / path):
            if cand.exists():
                return str(cand)
        raise FileNotFoundError(f"Imagem nao encontrada: {path}")

    @staticmethod
    def _set_alt(shape, text: str | None):
        if text:
            for el in shape._element.xpath(".//p:cNvPr"):
                el.set("descr", str(text))

    def add_image(self, title: str, image: str, caption: str | None = None, layout: str = "full", bullets: list[str] | None = None,
                  highlights: list[dict] | None = None, alt: str | None = None, subtitle: str | None = None, source: str | None = None,
                  notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        """Imagem/print ajustada a area de conteudo, com legenda, destaques numerados e texto alternativo."""
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        path = self._resolve(image)
        cap_h = 0.45 if caption else 0.0
        avail_h = self.bottom - top - 0.1 - cap_h
        if layout == "left":
            box, tb = (0.5, top + 0.1, 8.0, avail_h), (8.8, top + 0.1, 4.03, avail_h)
        elif layout == "right":
            box, tb = (4.83, top + 0.1, 8.0, avail_h), (0.5, top + 0.1, 4.03, avail_h)
        else:
            box, tb = (0.5, top + 0.1, CONTENT_W, avail_h), None
        try:
            from PIL import Image as _Img
            with _Img.open(path) as im:
                iw, ih = im.size
        except Exception:  # noqa: BLE001
            iw, ih = 16, 9
        bx, by, bw, bh = box
        scale = min(bw / iw, bh / ih)
        w, h = iw * scale, ih * scale
        x, y = bx + (bw - w) / 2, by + (bh - h) / 2
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        pic.line.color.rgb = self.c("neutral_light")
        pic.line.width = Pt(0.75)
        self._set_alt(pic, alt or caption or title)
        for i, hl in enumerate(highlights or []):
            hx, hy = x + float(hl.get("x", 0)) * w, y + float(hl.get("y", 0)) * h
            hw, hh = float(hl.get("w", 0.2)) * w, float(hl.get("h", 0.2)) * h
            self._rect(s, hx, hy, hw, hh, fill=None, line="accent", line_width=2.25)
            num = str(hl.get("number", i + 1))
            circ = self._rect(s, hx - 0.18, hy - 0.18, 0.36, 0.36, fill="accent", line="FFFFFF", shape=MSO_SHAPE.OVAL, line_width=1.0)
            circ.text_frame.text = num
            circ.text_frame.margin_left = circ.text_frame.margin_right = Inches(0)
            circ.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            circ.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._style_run(circ.text_frame.paragraphs[0].runs[0], 11, bold=True, color="FFFFFF")
        if caption:
            self._text(s, bx, by + bh + 0.05, bw, 0.4, caption, size=11, italic=True, color="neutral_mid", align="center")
        if tb and bullets:
            self._bullets(s, *tb, bullets, size=14)
        return s

    def _line(self, slide, x1, y1, x2, y2, color="neutral_mid", width=1.0, dash=False):
        ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = self.c(color)
        ln.line.width = Pt(width)
        if dash:
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return ln

    def add_waterfall(self, title: str, items: list[dict], subtitle: str | None = None, decimals: int = 1, unit: str = "",
                      source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        """Ponte (waterfall) com formas: items = [{label, value, total?}]; totais em primary, ganhos success, perdas danger."""
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        n = len(items)
        px, py, pw, ph = 0.5, top + 0.35, CONTENT_W, self.bottom - top - 1.15
        # cumulativo
        levels, cum = [], 0.0
        for it in items:
            v = float(it.get("value", 0))
            if it.get("total"):
                levels.append((0.0, v)); cum = v
            else:
                levels.append((cum, cum + v)); cum += v
        lo = min(0.0, min(min(a, b) for a, b in levels))
        hi = max(max(a, b) for a, b in levels)
        span = (hi - lo) or 1.0
        y_of = lambda v: py + ph - (v - lo) / span * ph  # noqa: E731
        slot = pw / n
        bw = slot * 0.62
        self._line(s, px, y_of(0), px + pw, y_of(0), color="neutral_mid", width=0.75)
        prev_top = None
        for i, (it, (a, b)) in enumerate(zip(items, levels)):
            x = px + i * slot + (slot - bw) / 2
            v = float(it.get("value", 0))
            color = "primary" if it.get("total") else ("success" if v >= 0 else "danger")
            y1, y2 = y_of(max(a, b)), y_of(min(a, b))
            self._rect(s, x, y1, bw, max(y2 - y1, 0.03), fill=color)
            label = (f"{'+' if (v > 0 and not it.get('total')) else ''}{fmt_number(v, decimals)}{(' ' + unit) if unit else ''}")
            self._text(s, x - 0.3, y1 - 0.38, bw + 0.6, 0.35, label, size=12, bold=True, color=color if not it.get("total") else "neutral_dark", align="center", anchor=MSO_ANCHOR.BOTTOM)
            self._text(s, px + i * slot, py + ph + 0.08, slot, 0.6, str(it.get("label", "")), size=11, color="neutral_dark", align="center")
            if prev_top is not None:
                self._line(s, x - (slot - bw), prev_top, x, prev_top, color="neutral_mid", width=0.75, dash=True)
            prev_top = y_of(b)
        return s

    def add_pareto(self, title: str, categories: list[str], values: list[float], subtitle: str | None = None, sort: bool = True,
                   threshold: float = 80.0, unit: str = "", source: str | None = None, notes: str | None = None,
                   kicker: str | None = None, callout: dict | None = None):
        """Pareto com formas: barras decrescentes + linha de % acumulado; barras ate o limiar em accent."""
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        pairs = list(zip(categories, [float(v) for v in values]))
        if sort:
            pairs.sort(key=lambda kv: -kv[1])
        total = sum(v for _, v in pairs) or 1.0
        n = len(pairs)
        px, py, pw, ph = 0.5, top + 0.4, CONTENT_W - 1.2, self.bottom - top - 1.3
        vmax = max(v for _, v in pairs) or 1.0
        slot = pw / n
        bw = slot * 0.6
        self._line(s, px, py + ph, px + pw, py + ph, color="neutral_mid", width=0.75)
        cum, pts = 0.0, []
        reached = False
        for i, (cat, v) in enumerate(pairs):
            x = px + i * slot + (slot - bw) / 2
            h = v / vmax * ph
            color = "accent" if not reached else "primary"
            self._rect(s, x, py + ph - h, bw, h, fill=color)
            self._text(s, x - 0.2, py + ph - h - 0.32, bw + 0.4, 0.3, fmt_number(v) + (f" {unit}" if unit else ""), size=11, bold=True, align="center", anchor=MSO_ANCHOR.BOTTOM)
            self._text(s, px + i * slot, py + ph + 0.06, slot, 0.7, cat, size=11, align="center")
            cum += v
            pct = 100 * cum / total
            pts.append((x + bw / 2, py + ph - pct / 100 * ph, pct))
            if pct >= threshold:
                reached = True
        for (x1, y1, _), (x2, y2, _) in zip(pts, pts[1:]):
            self._line(s, x1, y1, x2, y2, color="neutral_dark", width=1.75)
        for x, y, pct in pts:
            self._rect(s, x - 0.07, y - 0.07, 0.14, 0.14, fill="neutral_dark", shape=MSO_SHAPE.OVAL)
            self._text(s, x - 0.5, y - 0.4, 1.0, 0.3, f"{fmt_number(pct)}%", size=10, color="neutral_dark", align="center", anchor=MSO_ANCHOR.BOTTOM)
        ty = py + ph - threshold / 100 * ph
        self._line(s, px, ty, px + pw, ty, color="neutral_mid", width=0.75, dash=True)
        self._text(s, px + pw + 0.05, ty - 0.15, 1.1, 0.3, f"{fmt_number(threshold)}% acum.", size=10, color="neutral_mid", anchor=MSO_ANCHOR.MIDDLE)
        return s

    def add_gantt(self, title: str, periods: list[str], tasks: list[dict], subtitle: str | None = None, today: float | None = None,
                  source: str | None = None, notes: str | None = None, kicker: str | None = None, callout: dict | None = None):
        """Gantt simplificado: tasks = [{name, start, end, status?, milestone?}] com start/end em indice de periodo (fracao ok)."""
        s, top = self._content_slide(title, subtitle, source, notes, kicker, callout)
        name_w = 3.0
        gx, gy, gw = 0.5 + name_w, top + 0.15, CONTENT_W - name_w
        n_p = max(len(periods), 1)
        cw = gw / n_p
        head_h = 0.4
        row_h = min(0.5, (self.bottom - gy - head_h - 0.1) / max(len(tasks), 1))
        gh = head_h + row_h * len(tasks)
        self._rect(s, gx, gy, gw, head_h, fill="primary")
        for i, per in enumerate(periods):
            self._text(s, gx + i * cw, gy, cw, head_h, str(per), size=11, bold=True, color="FFFFFF", align="center", anchor=MSO_ANCHOR.MIDDLE)
        for j, tk in enumerate(tasks):
            y = gy + head_h + j * row_h
            if j % 2 == 1:
                self._rect(s, 0.5, y, CONTENT_W, row_h, fill="neutral_light")
            self._text(s, 0.55, y, name_w - 0.1, row_h, str(tk.get("name", "")), size=12, anchor=MSO_ANCHOR.MIDDLE)
            st = tk.get("status")
            color = st if st in ("success", "warning", "danger") else "primary"
            start = float(tk.get("start", 0))
            end = float(tk.get("end", start + 1))
            if tk.get("milestone"):
                mx = gx + start * cw
                self._rect(s, mx - 0.14, y + row_h / 2 - 0.14, 0.28, 0.28, fill=color, shape=MSO_SHAPE.DIAMOND)
            else:
                self._rect(s, gx + start * cw, y + row_h * 0.25, max((end - start) * cw, 0.05), row_h * 0.5, fill=color)
                if tk.get("label"):
                    self._text(s, gx + start * cw + 0.05, y, (end - start) * cw, row_h, str(tk["label"]), size=9, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
        for i in range(n_p + 1):
            self._line(s, gx + i * cw, gy + head_h, gx + i * cw, gy + gh, color="neutral_light", width=0.5)
        if today is not None:
            tx = gx + float(today) * cw
            self._line(s, tx, gy, tx, gy + gh, color="accent", width=1.5, dash=True)
            self._text(s, tx - 0.4, gy + gh + 0.02, 0.8, 0.3, "Hoje", size=10, bold=True, color="accent", align="center")
        return s

    def add_quote(self, text: str, author: str | None = None, dark: bool = False, notes: str | None = None):
        s = self._new_slide()
        if dark:
            self._fill_bg(s, "primary")
        txt_color = "FFFFFF" if dark else "primary"
        self._text(s, 1.0, 1.2, 1.2, 1.4, "“", size=110, bold=True, color="accent", font=self.font_title)
        self._text(s, 1.5, 2.2, 10.33, 2.6, text, size="quote", italic=True, color=txt_color, anchor=MSO_ANCHOR.MIDDLE)
        if author:
            self._text(s, 1.5, 5.0, 10.33, 0.5, f"— {author}", size=15, color="D9E1EA" if dark else "neutral_mid")
        self._notes(s, notes)
        return s

    def add_closing(self, title: str, subtitle: str | None = None, notes: str | None = None):
        if self.template_branding:
            s = self.prs.slides.add_slide(self._layouts["cover"])
            self.page += 1
            self._fill_placeholders(s, title, subtitle or "")
            self._notes(s, notes)
            return s
        s = self._new_slide()
        self._fill_bg(s, "primary_dark")
        self._text(s, 0.8, 2.6, 11.7, 1.4, title, size="closing", bold=True, color="FFFFFF", anchor=MSO_ANCHOR.BOTTOM, font=self.font_title)
        if subtitle:
            self._text(s, 0.8, 4.1, 11.7, 0.9, subtitle, size=19, color="D9E1EA")
        self._rect(s, 0, 6.9, SLIDE_W, 0.6, fill="accent")
        self._place_logo(s, 11.3, 0.5, width=1.5, dark=True)
        self._notes(s, notes)
        return s

    # ------------------------------------------------------------ spec API
    def add_from_spec(self, sl: dict):
        """Adiciona um slide a partir do dicionario de spec (campo 'type' define o modelo)."""
        t = sl.get("type")
        args = {k: v for k, v in sl.items() if k != "type"}
        if t == "progress_bars" and "max" in args:
            args["max_value"] = args.pop("max")
        dispatch = {
            "cover": self.add_cover, "agenda": self.add_agenda, "section": self.add_section,
            "executive_summary": self.add_executive_summary, "kpi_row": self.add_kpi_row, "big_number": self.add_big_number,
            "chart": self.add_chart, "table": self.add_table, "bullets": self.add_bullets, "two_column": self.add_two_column,
            "comparison": self.add_comparison, "timeline": self.add_timeline, "process": self.add_process,
            "matrix_2x2": self.add_matrix_2x2, "action_plan": self.add_action_plan, "quote": self.add_quote, "closing": self.add_closing,
            "takeaways": self.add_takeaways, "progress_bars": self.add_progress_bars, "image": self.add_image,
            "waterfall": self.add_waterfall, "pareto": self.add_pareto, "gantt": self.add_gantt,
        }
        if t not in dispatch:
            raise ValueError(f"Tipo de slide desconhecido: {t!r}. Tipos validos: {sorted(dispatch)}")
        return dispatch[t](**args)

    @classmethod
    def from_spec(cls, spec: dict, template: str | None = None) -> "DeckBuilder":
        meta = spec.get("meta", {})
        template = template or meta.get("template")
        if template and not Path(template).exists() and (ROOT / template).exists():
            template = str(ROOT / template)
        confidentiality = meta.get("confidentiality")
        if not confidentiality and meta.get("exposure"):
            confidentiality = {"interno": "Uso interno", "interareas": "Uso interno", "externo": "Confidencial"}.get(meta["exposure"])
        db = cls(palette=meta.get("palette", "corporativa-azul"), font_scale=meta.get("font_scale"),
                 confidentiality=confidentiality, logo=meta.get("logo"), template=template,
                 template_layouts=meta.get("template_layouts"), template_branding=meta.get("template_branding", False),
                 brand=meta.get("brand"), deck_name=meta.get("deck_name"), date=meta.get("date"),
                 logo_light=meta.get("logo_light"), logo_position=meta.get("logo_position", "footer"))
        for sl in spec.get("slides", []):
            db.add_from_spec(sl)
        return db

    def save(self, path: str | Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path
