"""
Extrai a estrutura de um .pptx existente para revisao (00-guia/revisao-de-deck-existente.md).

Saidas:
  - Outline em Markdown (padrao): por slide, titulo, textos (com tamanho de fonte), graficos (tipo,
    categorias, series), tabelas, imagens, notas, e diagnostico (titulo nao conclusivo, excesso de
    texto, sem fonte, muitos bullets).
  - --json: a mesma estrutura em JSON.
  - --to-spec: spec aproximada (schemas/deck-spec.schema.json) para reconstruir o deck no padrao do
    compendio: capa, bullets, graficos (com dados), tabelas, imagens, encerramento.

Uso:
    python extract_deck.py deck.pptx [--out outline.md] [--json outline.json] [--to-spec spec.json]
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
CONCLUSION_VERBS = re.compile(r"\b(é|são|está|estão|foi|foram|cresc|caiu|cai|reduz|aument|sub|concentr|respond|exig|precis|propo|recomend|deve|pode|mant|atras|supera|cobre|ganh|perd|explica|mostr|indica|sugere|ainda|apesar|mas|não|nao)\w*", re.I)
GENERIC_TITLES = {"agenda", "overview", "visão geral", "visao geral", "status", "dados", "análise", "analise", "resultados", "conclusão", "conclusao", "próximos passos", "proximos passos", "introdução", "introducao", "contexto", "obrigado", "sumário executivo", "sumario executivo", "riscos", "cronograma", "orçamento", "orcamento"}


def is_conclusion(title: str) -> tuple[bool, str]:
    t = title.strip()
    low = t.lower().rstrip(".:")
    if not t:
        return False, "sem título"
    if low in GENERIC_TITLES:
        return False, "título genérico (assunto, não conclusão)"
    words = t.split()
    if t.isupper() and len(words) > 1:
        return False, "caixa alta"
    if t.endswith(":"):
        return False, "termina com dois pontos"
    if len(words) < 4:
        return False, "menos de 4 palavras"
    if len(t) > 110:
        return False, "mais de ~2 linhas"
    if not (CONCLUSION_VERBS.search(t) or any(ch.isdigit() for ch in t)):
        return False, "sem verbo nem número (parece assunto)"
    return True, ""


def shape_text(sh) -> tuple[str, list[float]]:
    if not sh.has_text_frame:
        return "", []
    sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
    return sh.text_frame.text.strip(), sizes


def extract(path: Path) -> dict:
    prs = Presentation(str(path))
    deck = {"arquivo": path.name, "largura": round(prs.slide_width / EMU, 2), "altura": round(prs.slide_height / EMU, 2), "slides": []}
    for i, s in enumerate(prs.slides, start=1):
        info = {"n": i, "layout": s.slide_layout.name, "titulo": "", "textos": [], "graficos": [], "tabelas": [], "imagens": 0, "notas": "", "diagnostico": []}
        try:
            fill = s.background.fill
            info["fundo_escuro"] = bool(fill.type is not None and sum(fill.fore_color.rgb) < 300)
        except Exception:  # noqa: BLE001
            info["fundo_escuro"] = False
        for sh in s.shapes:
            if sh.is_placeholder and sh.placeholder_format.type is not None and "TITLE" in str(sh.placeholder_format.type):
                info["titulo"] = shape_text(sh)[0]
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                g = {"tipo": str(ch.chart_type).split(" ")[0], "categorias": [], "series": []}
                try:
                    plot = ch.plots[0]
                    g["categorias"] = [str(c) for c in plot.categories]
                    for ser in plot.series:
                        g["series"].append({"name": str(ser.name), "values": [None if v is None else float(v) for v in ser.values]})
                except Exception:  # noqa: BLE001
                    pass
                info["graficos"].append(g)
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                info["tabelas"].append(rows)
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                info["imagens"] += 1
                continue
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub in sh.shapes:
                    txt, sizes = shape_text(sub)
                    if txt:
                        info["textos"].append({"texto": txt, "fonte_min": min(sizes) if sizes else None, "top": round(sub.top / EMU, 2) if sub.top is not None else None})
                continue
            txt, sizes = shape_text(sh)
            if txt:
                info["textos"].append({"texto": txt, "fonte_min": min(sizes) if sizes else None, "top": round(sh.top / EMU, 2) if sh.top is not None else None})
        if not info["titulo"] and info["textos"]:
            # titulo provavel: maior fonte entre os textos do terco superior (desempate: mais alto)
            top_zone = [x for x in info["textos"] if x["top"] is not None and x["top"] < 2.3 and len(x["texto"].split()) >= 3] or \
                       [x for x in info["textos"] if x["top"] is not None and x["top"] < 2.3] or info["textos"]
            cand = sorted(top_zone, key=lambda x: (-(x["fonte_min"] or 0), x["top"] if x["top"] is not None else 9))[0]
            if (cand["fonte_min"] or 0) >= 18 or (cand["top"] is not None and cand["top"] < 1.2 and len(cand["texto"].split()) >= 4):
                info["titulo"] = cand["texto"].split("\n")[0]
                info["titulo_inferido"] = True
        if s.has_notes_slide:
            info["notas"] = s.notes_slide.notes_text_frame.text.strip()
        # diagnostico
        if not info["fundo_escuro"]:
            ok, why = is_conclusion(info["titulo"])
            if not ok:
                info["diagnostico"].append(f"título: {why}")
        body_words = sum(len(t["texto"].split()) for t in info["textos"])
        if body_words > 90:
            info["diagnostico"].append(f"texto denso ({body_words} palavras fora do título)")
        n_par = max([len([p for p in t["texto"].split("\n") if p.strip()]) for t in info["textos"]] or [0])
        if n_par > 7:
            info["diagnostico"].append(f"{n_par} linhas/bullets em uma caixa")
        has_data = bool(info["graficos"] or info["tabelas"])
        if has_data and not any("fonte" in t["texto"].lower()[:12] for t in info["textos"]):
            info["diagnostico"].append("gráfico/tabela sem fonte")
        small = [t["fonte_min"] for t in info["textos"] if t["fonte_min"] and t["fonte_min"] < 11 and (t["top"] or 0) < 6.6]
        if small:
            info["diagnostico"].append(f"fonte pequena no corpo ({min(small):.0f} pt)")
        if len(info["graficos"]) > 1:
            info["diagnostico"].append(f"{len(info['graficos'])} gráficos no mesmo slide")
        deck["slides"].append(info)
    return deck


def to_markdown(deck: dict) -> str:
    out = [f"# Outline: {deck['arquivo']}", "", f"Tamanho: {deck['largura']} x {deck['altura']} pol. | Slides: {len(deck['slides'])}", "",
           "## Teste dos títulos (a história se sustenta só com eles?)", ""]
    for s in deck["slides"]:
        flag = "" if not s["diagnostico"] else "  ⚠"
        out.append(f"{s['n']}. {s['titulo'] or '(sem título)'}{flag}")
    out += ["", "## Slide a slide", ""]
    for s in deck["slides"]:
        out += [f"### Slide {s['n']} - {s['titulo'] or '(sem título)'}", f"Layout: {s['layout']}" + (" | fundo escuro" if s["fundo_escuro"] else "") + (" | título inferido" if s.get("titulo_inferido") else ""), ""]
        for t in s["textos"]:
            fs = f" [{t['fonte_min']:.0f} pt]" if t["fonte_min"] else ""
            out.append(f"- {t['texto'][:220].replace(chr(10), ' | ')}{fs}")
        for g in s["graficos"]:
            out.append(f"- GRÁFICO {g['tipo']}: categorias {g['categorias'][:8]}; séries " + ", ".join(f"{x['name']} {x['values'][:8]}" for x in g["series"]))
        for tb in s["tabelas"]:
            out.append(f"- TABELA {len(tb)}x{len(tb[0]) if tb else 0}: cabeçalho {tb[0] if tb else []}")
        if s["imagens"]:
            out.append(f"- {s['imagens']} imagem(ns)")
        if s["notas"]:
            out.append(f"- NOTAS: {s['notas'][:200]}")
        if s["diagnostico"]:
            out.append("- DIAGNÓSTICO: " + "; ".join(s["diagnostico"]))
        out.append("")
    return "\n".join(out)


def to_spec(deck: dict) -> dict:
    slides = []
    first = deck["slides"][0] if deck["slides"] else None
    for s in deck["slides"]:
        title = s["titulo"] or "[TÍTULO-CONCLUSÃO]"
        body = [t["texto"] for t in s["textos"] if t["texto"] != s["titulo"] and (t["top"] or 0) < 6.6]
        if s is first and (s["fundo_escuro"] or "title" in s["layout"].lower() or "capa" in s["layout"].lower()):
            slides.append({"type": "cover", "title": title, "subtitle": body[0][:120] if body else None})
            continue
        if s["fundo_escuro"] and not s["graficos"] and not s["tabelas"] and len(body) <= 2:
            slides.append({"type": "section", "title": title, "subtitle": body[0][:120] if body else None})
            continue
        if s["graficos"]:
            g = s["graficos"][0]
            ctype = {"BAR": "bar", "COLUMN": "column", "LINE": "line", "PIE": "pie", "DOUGHNUT": "doughnut", "AREA": "area"}
            key = next((v for k, v in ctype.items() if g["tipo"].upper().startswith(k)), "column")
            if "STACKED" in g["tipo"].upper():
                key = "stacked_bar" if key == "bar" else "stacked_column"
            slides.append({"type": "chart", "title": title, "chart_type": key, "categories": g["categorias"], "series": [{"name": x["name"], "values": [v or 0 for v in x["values"]]} for x in g["series"]],
                           "commentary": [b[:120] for b in body[:3]] or None, "source": "[FONTE]"})
            continue
        if s["tabelas"]:
            tb = s["tabelas"][0]
            slides.append({"type": "table", "title": title, "columns": tb[0], "rows": tb[1:11], "source": "[FONTE]"})
            continue
        if s["imagens"]:
            slides.append({"type": "image", "title": title, "image": "[CAMINHO DA IMAGEM]", "bullets": [b[:120] for b in body[:5]] or None, "layout": "left" if body else "full"})
            continue
        bullets = []
        for b in body:
            bullets += [ln.strip() for ln in b.split("\n") if ln.strip()]
        slides.append({"type": "bullets", "title": title, "bullets": bullets[:6] or ["[CONTEÚDO]"]})
    if slides and slides[-1]["type"] != "closing":
        slides.append({"type": "closing", "title": "[PRÓXIMO PASSO / PEDIDO]"})
    for sl in slides:
        for k in [k for k, v in sl.items() if v is None]:
            del sl[k]
    return {"meta": {"title": deck["slides"][0]["titulo"] if deck["slides"] else deck["arquivo"], "audience": "tatico", "type": "[TIPO]", "palette": "corporativa-azul", "exposure": "interareas"}, "slides": slides}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--out")
    ap.add_argument("--json")
    ap.add_argument("--to-spec")
    a = ap.parse_args(argv)
    deck = extract(Path(a.deck))
    md = to_markdown(deck)
    out = Path(a.out) if a.out else Path(a.deck).with_suffix(".outline.md")
    out.write_text(md, encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    print(md)
    if a.json:
        Path(a.json).write_text(json.dumps(deck, ensure_ascii=False, indent=2), encoding="utf-8")
    if a.to_spec:
        Path(a.to_spec).write_text(json.dumps(to_spec(deck), ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\nSpec aproximada: {a.to_spec}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
